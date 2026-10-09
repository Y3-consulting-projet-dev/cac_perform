"""
Service d'aperçu de mapping pour l'import de balance générale.
Nouveau périmètre V2 — voir docs/PLAN_import_refonte.md §1a et §3.

Ne modifie ni n'importe api/src/model.py (CLAUDE.md : le parsing historique
reste isolé). Reprend, en lecture seule et de façon simplifiée, le même
principe de détection par défaut que Mission.creation_balance() : un format
connu est proposé selon le nombre de colonnes détectées. Cette proposition
sert de VALEUR PAR DÉFAUT que l'utilisateur confirme ou corrige avant tout
import définitif — elle ne remplace pas le parsing historique, qui continue
de tourner tel quel pour l'import actuel (balance générale N/N-1 à la
création de mission).
"""
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

NOMS_FEUILLES_CONNUES = [
    "Balance des comptes", "Balance_des_comptes", "BALANCE_2023", "BALANCE_2024",
    "Sage", "Sheet1", "Feuil1", "Balance", "Comptes",
]

# Formats connus : nombre de colonnes -> (nom du format, champs cibles dans l'ordre).
# Même inventaire que Mission.creation_balance() (api/src/model.py), reconstruit
# indépendamment pour ne pas coupler ce service au fichier historique.
FORMATS_CONNUS: Dict[int, Tuple[str, List[str]]] = {
    8: ("standard", [
        "numero", "libelle",
        "debit_initial", "credit_initial",
        "debit_mvt", "credit_mvt",
        "debit_fin", "credit_fin",
    ]),
    6: ("balance_simple", [
        "numero", "libelle",
        "debit_initial", "credit_initial",
        "debit_fin", "credit_fin",
    ]),
    3: ("simple", ["numero", "libelle", "montant"]),
}
# Remarque : le format "sage" (6 colonnes) partage son nombre de colonnes avec
# "balance_simple" ; model.py les distingue via une heuristique plus fine sur
# le contenu des en-têtes, non reproduite ici.
# TODO(spec: confirmer si l'aperçu V2 doit distinguer sage/balance_simple
# automatiquement, ou si la correction manuelle par l'utilisateur suffit).


def _choisir_feuille(workbook):
    for nom in NOMS_FEUILLES_CONNUES:
        if nom in workbook.sheetnames:
            return workbook[nom]
    return workbook[workbook.sheetnames[0]]


def generer_apercu_mapping(fichier) -> Dict[str, Any]:
    """
    Lit un fichier Excel et propose un mapping de colonnes par défaut.
    Ne modifie rien en base. `fichier` : objet compatible
    openpyxl.load_workbook (ex. werkzeug FileStorage).
    """
    workbook = openpyxl.load_workbook(fichier, data_only=True)
    sheet = _choisir_feuille(workbook)

    premiere_ligne = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    # Les fichiers réels ont souvent des colonnes vides en fin de ligne
    # (mise en forme résiduelle) : on ne les compte pas comme des colonnes
    # de données, sinon le nombre de colonnes ne correspond plus à aucun
    # format connu.
    while premiere_ligne and premiere_ligne[-1] is None:
        premiere_ligne.pop()

    colonnes_detectees = [
        str(c).strip() if c is not None else f"Colonne {i + 1}"
        for i, c in enumerate(premiere_ligne)
    ]

    nb_lignes_detectees = max(sheet.max_row - 1, 0)
    nb_colonnes = len(colonnes_detectees)

    format_connu: Optional[Tuple[str, List[str]]] = FORMATS_CONNUS.get(nb_colonnes)
    if format_connu:
        format_propose, champs = format_connu
    else:
        format_propose, champs = "inconnu", []

    mapping_propose: Dict[str, str] = {}
    for index, champ in enumerate(champs):
        if index < len(colonnes_detectees):
            mapping_propose[champ] = colonnes_detectees[index]

    return {
        "colonnes_detectees": colonnes_detectees,
        "feuille": sheet.title,
        "format_propose": format_propose,
        "mapping_propose": mapping_propose,
        "nb_lignes_detectees": nb_lignes_detectees,
    }
