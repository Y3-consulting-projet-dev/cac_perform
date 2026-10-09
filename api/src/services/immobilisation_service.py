"""
Service d'import du fichier des immobilisations.
Nouveau périmètre V2 — voir docs/PLAN_import_refonte.md §1c et §2.

Plomberie uniquement : réception, parsing générique, stockage. AUCUNE règle
de rapprochement IPE-3 ici (le verrou IPE/Module 1 est hors périmètre de ce
chantier, voir CLAUDE.md). Ne modifie pas api/src/model.py.

TODO(spec: CAC-IPE-01 §3) — colonnes exactes attendues (valeur brute,
amortissements cumulés, valeur nette comptable, durée/taux d'amortissement,
compte 2xxx concerné...) non confirmées dans le dépôt. Parsing générique en
attendant la spec.
"""
from typing import Any, Dict, List

import openpyxl

from src.utils.database import get_database


def _parser_generique(fichier) -> List[Dict[str, Any]]:
    """Lecture générique à 3 colonnes : numéro, libellé, montant.
    TODO(spec: CAC-IPE-01 §3) — structure réelle à confirmer, voir en-tête du module."""
    workbook = openpyxl.load_workbook(fichier, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    lignes = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        numero = row[0] if len(row) > 0 else None
        libelle = row[1] if len(row) > 1 else None
        montant = row[2] if len(row) > 2 else None
        if numero is None:
            continue
        lignes.append({
            "numero": str(numero).strip(),
            "libelle": str(libelle).strip() if libelle is not None else "",
            "montant": float(montant) if isinstance(montant, (int, float)) else 0.0,
        })
    return lignes


def importer_immobilisations(fichier, id_client: str, mission_id: str, annee: str) -> Dict[str, Any]:
    """
    Importe le fichier des immobilisations et le stocke dans la collection
    Balance avec `type="immobilisations"` + `mission_id` (voir
    docs/PLAN_import_refonte.md §2).
    """
    lignes = _parser_generique(fichier)
    if not lignes:
        raise ValueError("Aucune ligne exploitable trouvée dans le fichier")

    db = get_database()
    result = db.Balance.insert_one({
        "id_client": id_client,
        "mission_id": mission_id,
        "annee_balance": annee,
        "type": "immobilisations",
        "balance": lignes,
    })
    return {
        "_id": str(result.inserted_id),
        "type": "immobilisations",
        "nb_lignes": len(lignes),
    }
