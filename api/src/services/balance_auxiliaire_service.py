"""
Service d'import des balances auxiliaires (clients / fournisseurs).
Nouveau périmètre V2 — voir docs/PLAN_import_refonte.md §1b et §2.

Plomberie uniquement : réception, parsing générique, stockage. AUCUNE règle
de rapprochement IPE ici (le verrou IPE/Module 1 est hors périmètre de ce
chantier, voir CLAUDE.md). Ne modifie pas api/src/model.py ni les documents
Balance de type "generale" existants.

TODO(spec: CAC-IPE-01) — les colonnes exactes attendues pour une balance
auxiliaire (ex. compte collectif tiers, code tiers, raison sociale, solde
débiteur/créditeur...) ne sont pas confirmées dans le dépôt. Le parsing
ci-dessous reste générique (numéro de tiers, libellé, montant) en attendant
la spec ; à affiner une fois CAC-IPE-01 confirmée.
"""
from typing import Any, Dict, List

import openpyxl

from src.utils.database import get_database

TYPES_VALIDES = ("client", "fournisseur")


def _parser_generique(fichier) -> List[Dict[str, Any]]:
    """Lecture générique à 3 colonnes : numéro de tiers, libellé, montant.
    TODO(spec: CAC-IPE-01) — structure réelle à confirmer, voir en-tête du module."""
    workbook = openpyxl.load_workbook(fichier, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    lignes = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        numero_tiers = row[0] if len(row) > 0 else None
        libelle = row[1] if len(row) > 1 else None
        montant = row[2] if len(row) > 2 else None
        if numero_tiers is None:
            continue
        lignes.append({
            "numero_tiers": str(numero_tiers).strip(),
            "libelle": str(libelle).strip() if libelle is not None else "",
            "montant": float(montant) if isinstance(montant, (int, float)) else 0.0,
        })
    return lignes


def importer_balance_auxiliaire(fichier, type_auxiliaire: str, id_client: str, mission_id: str, annee: str) -> Dict[str, Any]:
    """
    Importe une balance auxiliaire (clients ou fournisseurs) et la stocke
    dans la collection Balance avec un champ `type` dédié + `mission_id`
    (voir docs/PLAN_import_refonte.md §2). N'altère aucun document Balance
    existant de type "generale".
    """
    if type_auxiliaire not in TYPES_VALIDES:
        raise ValueError(f"type_auxiliaire doit être 'client' ou 'fournisseur' (reçu: '{type_auxiliaire}')")

    lignes = _parser_generique(fichier)
    if not lignes:
        raise ValueError("Aucune ligne exploitable trouvée dans le fichier")

    db = get_database()
    result = db.Balance.insert_one({
        "id_client": id_client,
        "mission_id": mission_id,
        "annee_balance": annee,
        "type": f"auxiliaire_{type_auxiliaire}",
        "balance": lignes,
    })
    return {
        "_id": str(result.inserted_id),
        "type": f"auxiliaire_{type_auxiliaire}",
        "nb_lignes": len(lignes),
    }
