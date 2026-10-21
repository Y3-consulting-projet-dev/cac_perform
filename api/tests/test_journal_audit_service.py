"""
Journal d'audit (page B) : assemblage en lecture seule de l'historique déjà
enregistré, filtres, pagination, export Excel. Données FICTIVES en mémoire.
"""
import inspect
import io
from datetime import date

import openpyxl

from src.services import journal_audit_service as journal


def _mission():
    return {
        "balance_generale_remplacements": [{
            "horodatage": "2026-10-13T09:00:00",
            "utilisateur_id": "Auditeur Test",
            "version_remplacee": 1,
            "nouvelle_version": 2,
            "nom_fichier": "bg_fictive.xlsx",
            "decisions_sous_totaux": [{"numero_compte": "63", "libelle": "SERVICES EXTERIEURS", "decision": "ecarter"}],
        }],
        "ipe": {
            "remarques": [{"_id": "r1", "commentaire": "Vérifier le compte 466", "module": "ipe_2_fournisseurs",
                           "numero_compte": "401"}],
            "historique": [
                {"horodatage": "2026-10-15T16:37:00+00:00", "utilisateur_id": None, "action": "evaluation",
                 "resultat": {"ipe_1_clients": "ecart", "ipe_2_fournisseurs": "ok", "ipe_3_immobilisations": "non_evaluable"},
                 "commentaire": None},
                {"horodatage": "2026-10-15T17:00:00+00:00", "utilisateur_id": "Manager Test", "action": "validation_manager",
                 "resultat": {"validee": True}, "commentaire": "Rapprochement revu et conforme"},
                {"horodatage": "2026-10-15T16:50:00+00:00", "utilisateur_id": "Auditeur Test", "action": "remarque_traitee",
                 "resultat": {"remarque_id": "r1"}, "commentaire": None},
                {"horodatage": "2026-10-15T16:55:00+00:00", "utilisateur_id": "Auditeur Test", "action": "action_future",
                 "resultat": {}, "commentaire": None},
            ],
        },
    }


def _documents():
    return [
        {"type": "auxiliaire_fournisseur", "nom_fichier": "aux_frs.xlsx", "version": 1, "date_import": "2026-10-15T16:12:00",
         "balance": [{}, {}, {}],
         "retrait": {"horodatage": "2026-10-16T08:00:00+00:00", "utilisateur_id": "Auditeur Test",
                     "commentaire": "je n'en ai pas besoin"}},
        # La balance générale V2 est tracée par balance_generale_remplacements : pas de doublon.
        {"type": "balance_generale", "nom_fichier": "bg_fictive.xlsx", "version": 2, "date_import": "2026-10-13T09:00:00",
         "balance": [{}]},
    ]


def _journal():
    return journal.assembler(_mission(), _documents())


def test_assemblage_trie_du_plus_recent_au_plus_ancien():
    entrees = _journal()
    dates = [e["horodatage"] for e in entrees]
    assert dates == sorted(dates, reverse=True)
    assert entrees[0]["action"] == "Fichier retiré"
    assert entrees[0]["detail"] == "aux_frs.xlsx · v1 · « je n'en ai pas besoin »"


def test_import_sans_auteur_reste_non_identifie_et_bg_pas_comptee_deux_fois():
    imports = [e for e in _journal() if e["action"] == "Fichier importé"]
    assert len(imports) == 1
    assert imports[0]["utilisateur"] is None
    assert imports[0]["detail"] == "Balance auxiliaire fournisseurs · aux_frs.xlsx · v1 · 3 lignes"
    assert imports[0]["horodatage"] == "2026-10-15T16:12:00+00:00"  # sans fuseau = UTC


def test_remplacement_et_decision_sous_total():
    entrees = {e["action"]: e for e in _journal()}
    assert entrees["Balance générale remplacée"]["detail"] == "bg_fictive.xlsx · version 1 → 2"
    assert entrees["Décision sur un sous-total"]["detail"] == "compte 63 · SERVICES EXTERIEURS · ligne écartée"
    assert entrees["Décision sur un sous-total"]["module"] == "Import"


def test_historique_ipe():
    entrees = {e["action"]: e for e in _journal()}
    assert entrees["Rapprochement lancé"]["detail"] == "IPE-1 : écart · IPE-2 : OK · IPE-3 : non évaluable"
    assert entrees["Validation"]["detail"] == "Commentaire « Rapprochement revu et conforme »"
    assert entrees["Remarque marquée traitée"]["detail"] == "IPE-2 · compte 401 · remarque « Vérifier le compte 466 »"
    # Une action inconnue du journal est affichée telle quelle, jamais masquée.
    assert "action_future" in entrees


def test_filtres_module_utilisateur_et_periode():
    entrees = _journal()
    assert {e["module"] for e in journal.filtrer(entrees, module="Fiabilité IPE")} == {"Fiabilité IPE"}
    non_identifies = journal.filtrer(entrees, utilisateur=journal.FILTRE_NON_IDENTIFIE)
    assert {e["action"] for e in non_identifies} == {"Rapprochement lancé", "Fichier importé"}
    assert {e["utilisateur"] for e in journal.filtrer(entrees, utilisateur="Manager Test")} == {"Manager Test"}
    du_15 = journal.filtrer(entrees, du=date(2026, 10, 15), au=date(2026, 10, 15))
    assert len(du_15) == 5
    # 16/10 08:00 UTC = 15/10 23:00 en UTC-9 : la période est lue dans l'heure de l'utilisateur.
    assert any(e["action"] == "Fichier retiré"
               for e in journal.filtrer(entrees, au=date(2026, 10, 15), decalage_minutes=-540))


def test_options_utilisateurs():
    options = journal.utilisateurs(_journal())
    assert [o["libelle"] for o in options] == ["Auditeur Test", "Manager Test", "Non identifié"]
    assert options[-1]["valeur"] == journal.FILTRE_NON_IDENTIFIE


def test_pagination():
    entrees = [{"horodatage": None}] * 30
    page = journal.paginer(entrees, 2, taille=25)
    assert (page["page"], page["nb_pages"], page["nb_entrees"], len(page["entrees"])) == (2, 2, 30, 5)
    assert journal.paginer([], 5)["page"] == 1
    assert journal.paginer(entrees, 99, taille=25)["page"] == 2


def test_export_excel():
    contenu = journal.exporter_excel(_journal(), decalage_minutes=60)
    ws = openpyxl.load_workbook(io.BytesIO(contenu)).active
    assert [c.value for c in ws[1]] == list(journal.COLONNES)
    assert ws.max_row == len(_journal()) + 1
    premiere = [c.value for c in ws[2]]
    assert premiere[0].strftime("%d/%m/%Y %H:%M") == "16/10/2026 09:00"  # heure de l'utilisateur
    assert premiere[1:4] == ["Auditeur Test", "Import", "Fichier retiré"]
    assert "Non identifié" in [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]


def test_le_service_n_ecrit_jamais_en_base():
    source = inspect.getsource(journal)
    for ecriture in ("insert_", "update_", "delete_", "replace_one", "find_one_and", "bulk_write", "drop"):
        assert ecriture not in source
